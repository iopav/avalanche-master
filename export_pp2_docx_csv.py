"""Export one reference-styled DOCX and one task CSV per PP2 dataset.

Python 3.10+, python-docx. Reads existing results only; no Torch or training.
CSV includes all successful LR candidates and the existing paired joint runs.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import hashlib
import json
import math
from pathlib import Path
import runpy
import tempfile
import os
import uuid
from zipfile import ZipFile
from io import BytesIO

from generate_pp2_tables import DISPLAY, require, close, identity, matrix
from generate_pp2_dataset_tables import failure_cost
from export_tagfex_task_flops import candidate_rows

FIELDS = ('task', 'test_accuracy', 'pertaskflops', 'model_parameter_storage_bytes',
          'model_parameter_storage_bytes_corrected', 'param_memory_kib',
          'method', 'seed', 'order', 'run_id')
METRICS = ('average_incremental_accuracy', 'final_average_accuracy', 'average_forgetting',
           'mean_final_accuracy_loss', 'backward_transfer')


def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def integer(value):
    require(type(value) is int and value >= 0, f'Invalid nonnegative integer: {value}')
    return value


def backbone_parameters(config):
    """Exact registered model parameter counts; excludes BatchNorm buffers."""
    b = config['backbone']
    require(b['backbone_id'] == 'resnet18_cifar' and b['feature_dim'] == 8*b['base_width'],
            'Unverified backbone for task parameter reconstruction')
    require(b['stage_channels'] == [b['base_width'] * 2**i for i in range(4)], 'Backbone stages differ')
    channels = 3  # Both registered stored RGB views and the Spike adapter.
    width = b['base_width']
    total, previous = channels * width * 9 + 2 * width, width
    for current in b['stage_channels']:
        for _ in range(2):
            total += previous * current * 9 + current * current * 9 + 4 * current
            if previous != current:
                total += previous * current + 2 * current
            previous = current
    return total


def parameter_bytes(config, groups, method):
    """FP32 shapes reconstructed from models.py, strategies.py, tagfex_avalanche.py.

    IncrementalClassifier expands to cumulative remapped classes (also FeCAM's
    retained train classifier). TagFex retains TA, all TS branches, projection,
    classifiers, predictor and attention. Old teacher copies are auxiliary.
    """
    base, d, seen, result = backbone_parameters(config), config['backbone']['feature_dim'], 0, []
    require(method in (*DISPLAY, 'joint'), f'Unsupported model: {method}')
    for task, group in enumerate(groups, 1):
        seen += len(group)
        count = base + seen * (d + 1)
        if method == 'tagfex':
            hp = config['method_parameters']
            hidden, output = hp['proj_hidden_dim'], hp['proj_output_dim']
            count = (task + 1)*base + hidden*(d+1) + output*(hidden+1) + seen*(task*d+1)
            if task > 1:
                count += (2*len(group)+1)*(d+1) + d*(d+1) + 5*d*d + 4*d
        result.append(count * 4)
    return result


def performance(payload, counts):
    rows = matrix(payload, len(counts))
    require(all(type(n) is int and n > 0 for n in counts), 'Invalid test sample weights')
    curve = [math.fsum(row[j]*counts[j] for j in range(t+1))/sum(counts[:t+1])
             for t, row in enumerate(rows)]
    losses = [rows[i][i] - rows[-1][i] for i in range(len(rows)-1)]
    forget = [max(rows[j][i] for j in range(i, len(rows))) - rows[-1][i]
              for i in range(len(rows)-1)]
    require(losses, 'At least two tasks required for forgetting')
    metrics = dict(zip(METRICS, (math.fsum(curve)/len(curve), curve[-1],
                   math.fsum(forget)/len(forget), math.fsum(losses)/len(losses),
                   -math.fsum(losses)/len(losses))))
    return metrics, curve


def task_rows(method, reference, lr, order, seed, curve, params, flops):
    token = str(lr).replace('.', '')
    prefix = method if method != 'joint' else f'joint_for_{reference}'
    run_id = f'{prefix}_lr{token}_o{order}_s{seed}'
    require(len(curve) == len(params) == len(flops), 'Task row coverage mismatch')
    return [dict(zip(FIELDS, (t, accuracy, cost, size, size, size/1024,
                             method, seed, order, run_id)))
            for t, (accuracy, size, cost) in enumerate(zip(curve, params, flops), 1)]


def validate_ops(ops):
    core = integer(ops['core_training_flops'])
    aux = integer(ops['learning_auxiliary_flops'])
    require(integer(ops['overall_learning_flops']) == core + aux, 'Learning FLOPs sum mismatch')


def load_dataset(root, dataset, exp, orders, seeds, lrs, methods):
    search_root = root / f'search_result_{exp}'
    require({p.name for p in search_root.iterdir() if p.is_dir() and p.name in DISPLAY} == set(methods),
            'Missing method directory')
    runs, selected, joints, csv_rows, failures, notes = [], [], [], [], [], []
    for method in methods:
        for order, groups in orders.items():
            for seed in seeds:
                directory = search_root / method / f'order{order}'
                search = read(directory / f'order{order}_seed{seed:03d}_search.json')
                expected = dict(dataset=dataset, exp_name=exp, method=method, order=order, seed=seed)
                identity(search, dict(expected, status='completed'))
                require(search['lr_candidates'] == list(lrs), 'LR registry mismatch')
                require(set(search['candidates']) == {str(lr) for lr in lrs}, 'Candidate keys mismatch')
                best = search['best_lr']
                require(best in lrs, 'Unknown best LR')
                best_run = None
                known = 0
                has_failure = False
                for lr in lrs:
                    candidate = search['candidates'][str(lr)]
                    basic = dict(method=method, order=order, seed=seed, lr=lr)
                    if candidate['status'] == 'failed':
                        has_failure = True
                        failure = dict(basic, status='failed', error=candidate.get('error_message', 'Failed'))
                        runs.append(failure)
                        failures.append(failure)
                        continue
                    require(candidate['status'] == 'completed', 'Pending candidate')
                    token = str(lr).replace('.', '')
                    stem = f'search_{dataset}_{method}__order-{order}__lr-{token}__seed-{seed:03d}'
                    path = directory / f'lr{token}' / f'{stem}__summary.json'
                    summary = read(path)
                    config = summary['config']
                    identity(summary, dict(exp_name=exp, seed=seed, tasks=len(groups)))
                    identity(config, dict(exp_name=exp, method=method))
                    require(config['dataset']['name'] == dataset, 'Config dataset mismatch')
                    require(config['selected_task_groups'] == [list(g) for g in groups], 'Task groups differ from registry')
                    close(config['final_hyperparameters']['resolved']['learning_rate'], lr)
                    payload = read(path.with_name(f'{stem}__accuracy-matrix.json'))
                    identity(payload, dict(dataset=dataset, exp_name=exp, method=method, order_id=order, seed=seed))
                    metrics, curve = performance(payload, payload['test_samples_per_task'])
                    for key, value in metrics.items():
                        close(value, summary['cil_performance'][key])
                    require(len(summary['cil_performance']['task_end_seen_accuracy_curve']) == len(curve), 'Curve length mismatch')
                    for a, b in zip(curve, summary['cil_performance']['task_end_seen_accuracy_curve']):
                        close(a, b)
                    ops, storage = summary['training_operations'], summary['persistent_storage']
                    validate_ops(ops)
                    require(integer(candidate['overall_learning_flops']) == ops['overall_learning_flops'], 'Search/summary FLOPs mismatch')
                    close(candidate['selection_loss'], summary['selection']['loss'])
                    params = parameter_bytes(config, groups, method)
                    require(params[-1] == integer(storage['model_parameter_bytes']), f'Parameter reconstruction differs: {path}')
                    require(integer(storage['total_bytes']) == sum(integer(storage[k]) for k in
                            ('model_parameter_bytes','auxiliary_bytes','replay_sample_bytes','replay_label_bytes')), 'Storage sum mismatch')
                    flops = [r['task_learning_flops'] for r in candidate_rows(path, exp)] if method == 'tagfex' else ['']*len(groups)
                    csv_rows.extend(task_rows(method, method, lr, order, seed, curve, params, flops))
                    run = dict(basic, status='completed', metrics=metrics, ops=ops, storage=storage,
                               config=config, counts=payload['test_samples_per_task'],
                               accuracy_matrix=matrix(payload, len(groups)))
                    runs.append(run)
                    known += ops['overall_learning_flops']
                    if lr == best:
                        best_run = run
                require(best_run is not None, 'Selected candidate missing')
                if has_failure:
                    require(search.get('total_search_flops') is None, 'Failed search claims exact cost')
                    require(integer(search['search_flops_lower_bound']) == known, 'Search lower bound mismatch')
                else:
                    require(integer(search['total_search_flops']) == known, 'Search total mismatch')
                partial, evidence = failure_cost(directory, expected, len(groups))
                notes.extend(f'{DISPLAY[method]} {text}' for text in evidence)
                complete = not has_failure and not evidence and not any(c.get('previous_failure_file') for c in search['candidates'].values())
                selected.append(dict(best_run, search_extra=known+partial-best_run['ops']['overall_learning_flops'],
                                     total=known+partial, complete=complete))
                joint_path = root / f'joint_result_{exp}' / method / f'order{order}' / f'joint_{dataset}_{method}__order-{order}__seed-{seed:03d}.json'
                joint = read(joint_path)
                identity(joint, dict(expected, status='completed', best_lr=best, tasks=len(groups)))
                require(joint['task_groups'] == [list(g) for g in groups], 'Joint task groups differ')
                metrics, curve = performance(joint, best_run['counts'])
                joint_matrix = matrix(joint, len(groups))
                for run in runs:
                    if (run['method'], run['order'], run['seed'], run['status']) == (method, order, seed, 'completed'):
                        run['metrics']['intransigence_mean'] = math.fsum(
                            joint_matrix[t][t] - run['accuracy_matrix'][t][t]
                            for t in range(len(groups))) / len(groups)
                ops, original_storage = joint['training_operations'], joint['persistent_storage']
                validate_ops(ops)
                params = parameter_bytes(best_run['config'], groups, 'joint')
                require(params[-1] == integer(original_storage['model_parameter_bytes']), 'Joint parameter shape mismatch')
                auxiliary = integer(original_storage['model_buffer_bytes']) + integer(original_storage['optimizer_state_bytes'])
                require(original_storage['total_bytes'] == params[-1] + auxiliary, 'Joint storage sum mismatch')
                storage = dict(model_parameter_bytes=params[-1], auxiliary_bytes=auxiliary,
                               replay_sample_bytes=0, replay_label_bytes=0, total_bytes=original_storage['total_bytes'])
                joints.append(dict(method=method, order=order, seed=seed, lr=best, metrics=metrics,
                                   ops=ops, storage=storage, status='completed'))
                csv_rows.extend(task_rows('joint', method, best, order, seed, curve, params, ['']*len(groups)))
    require(len({(r['run_id'],r['task']) for r in csv_rows}) == len(csv_rows), 'Duplicate CSV row identity')
    return dict(dataset=dataset, exp=exp, runs=runs, selected=selected, joints=joints,
                csv_rows=csv_rows, failures=failures, notes=notes, methods=methods, lrs=lrs)


def make_sections(data):
    sections = []
    headers = ['O','S','AIA','FAA','AF','MFAL','BWT','Intrans.','Total train\nFLOPs\nE12','Core\nFLOPs\nE12',
               'Aux\nFLOPs\nE12','Para\nMiB','Aux\nMiB','Replay\nSamp\nMiB','Replay\nLabels\nKiB','Persist\nMiB']
    def cells(run):
        if run['status'] == 'failed':
            return [run['order'],run['seed'],'FAIL'] + ['—']*(len(headers)-3)
        result = [run['order'],run['seed']] + [f'{run["metrics"][k]:.4f}' for k in METRICS]
        result += [f'{run["metrics"]["intransigence_mean"]:.4f}' if 'intransigence_mean' in run['metrics'] else '—']
        result += [f'{run["ops"][k]/1e12:.3f}' for k in ('overall_learning_flops','core_training_flops','learning_auxiliary_flops')]
        result += [f'{run["storage"][k]/scale:.3f}' for k,scale in
                   [('model_parameter_bytes',2**20),('auxiliary_bytes',2**20),('replay_sample_bytes',2**20),
                    ('replay_label_bytes',1024),('total_bytes',2**20)]]
        return result
    for method in data['methods']:
        name = DISPLAY[method]
        for lr in data['lrs']:
            runs = [r for r in data['runs'] if r['method']==method and r['lr']==lr]
            sections.append((f'{name} LR {lr}', headers, [cells(r) for r in runs]))
        rows = []
        for run in data['selected']:
            if run['method'] != method:
                continue
            row = cells(run)
            prefix = '' if run['complete'] else '≥ '
            rows.append(row[:2]+[str(run['lr'])]+row[2:11]+
                        [prefix+f'{run["search_extra"]/1e12:.3f}',prefix+f'{run["total"]/1e12:.3f}']+row[11:])
        sections.append((f'{name} selected LR and search cost', headers[:2]+['Best\nLR']+headers[2:11]+
                         ['Search\nextra\nE12','Total learn\nFLOPs\nE12']+headers[11:], rows))
    for method in data['methods']:
        rows = []
        for run in data['joints']:
            if run['method'] == method:
                row = cells(run)
                rows.append(row[:2]+[str(run['lr'])]+row[2:])
        sections.append((f'Joint reference for {DISPLAY[method]}', headers[:2]+['LR']+headers[2:], rows))
    return sections


def write_docx(template, output, data):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.table import Table
    doc = Document(template)
    table_model = deepcopy(doc.tables[0]._tbl)
    body = doc._element.body
    for element in list(body):
        if element.tag != qn('w:sectPr'):
            body.remove(element)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.size = Pt(10 if name=='Normal' else 14 if name!='Title' else 18)
    doc.add_paragraph(f'{data["dataset"].capitalize()} experimental results', 'Title')
    p=doc.add_paragraph(f'Experiment {data["exp"]}    LR candidates 0.1  0.05  0.01')
    p.paragraph_format.space_after=Pt(4)
    sections = make_sections(data)
    doc.add_paragraph('Table of contents','Heading 1')
    for i,(title,_,_) in enumerate(sections,1):
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'table_{i}')
        run=OxmlElement('w:r');prop=OxmlElement('w:rPr')
        size=OxmlElement('w:sz');size.set(qn('w:val'),'20');prop.append(size)
        color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');prop.append(color)
        run.append(prop);text=OxmlElement('w:t');text.text=f'Table {i}  {title}';run.append(text);link.append(run);p._p.append(link)
    for i,(title,headers,rows) in enumerate(sections,1):
        if i%2==1:
            doc.add_page_break()
        p=doc.add_paragraph(f'Table {i}  {title}', 'Heading 1')
        p.paragraph_format.space_before=Pt(8 if i%2==0 else 0)
        p.paragraph_format.space_after=Pt(4)
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(i));start.set(qn('w:name'),f'table_{i}')
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i));p._p.insert(0,start);p._p.append(end)
        # Clone the actual reference table and adapt only its grid, rows and values.
        element=deepcopy(table_model)
        source_rows=element.findall(qn('w:tr'))
        header_model=deepcopy(source_rows[0]);body_model=deepcopy(source_rows[1])
        for row in source_rows:element.remove(row)
        grid=element.find(qn('w:tblGrid'))
        for child in list(grid):grid.remove(child)
        widths=[.28,.32]+[.52]*6+[.90,.82,.72]+[.72,.72,.72,.72,.78]
        if len(headers)==19:widths=[.28,.32,.40]+[.48]*5+[.56]+[.85,.80,.66,.82,.88]+[.58]*5
        elif len(headers)==17:widths=[.28,.32,.42]+widths[2:]
        available=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
        widths=[int(available*w/sum(widths)) for w in widths]
        for width in widths:
            column=OxmlElement('w:gridCol');column.set(qn('w:w'),str(round(width/635)));grid.append(column)
        for ri,values in enumerate([headers,*rows]):
            require(len(values)==len(headers),'DOCX cell count mismatch')
            row=deepcopy(header_model if ri==0 else body_model)
            cell_model=deepcopy(row.findall(qn('w:tc'))[0])
            for cell in row.findall(qn('w:tc')):row.remove(cell)
            for width,value in zip(widths,values):
                cell=deepcopy(cell_model)
                for child in list(cell):
                    if child.tag!=qn('w:tcPr'):cell.remove(child)
                tcw=cell.find(qn('w:tcPr')).find(qn('w:tcW'));tcw.set(qn('w:w'),str(round(width/635)))
                paragraph=OxmlElement('w:p');cell.append(paragraph)
                row.append(cell)
            element.append(row)
        doc._element.body.insert(len(doc._element.body)-1,element)
        table=Table(element,doc._body);table.autofit=False
        for ri,row in enumerate(table.rows):
            props=row._tr.get_or_add_trPr()
            for old in list(props):
                if old.tag in (qn('w:trHeight'),qn('w:cantSplit')):props.remove(old)
            props.append(OxmlElement('w:cantSplit'))
            if ri==0 and props.find(qn('w:tblHeader')) is None:props.append(OxmlElement('w:tblHeader'))
            for cell,value in zip(row.cells,[headers,*rows][ri]):
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=Pt(9);p.paragraph_format.keep_with_next=False
                run=p.add_run(str(value));run.font.name='Times New Roman';run.font.size=Pt(8);run.bold=ri==0
    doc.add_page_break()
    doc.add_paragraph('Table notes','Heading 1')
    definitions = [
        'O is order and S is seed. Each row is one run. Accuracy values are fractions, not percentages. AIA is the mean of sample-weighted seen-test accuracy over stages; FAA is the last stage value. AF is average forgetting; MFAL is mean final accuracy loss; BWT equals minus MFAL.',
        'Total train FLOPs = Core FLOPs + Aux FLOPs for one LR candidate. Total learn FLOPs = Total train FLOPs (LR = 0.1) + Total train FLOPs (LR = 0.05) + Total train FLOPs (LR = 0.01). Search extra = Total learn FLOPs minus the selected candidate Total train FLOPs. The selected candidate is counted once. E12 means 10^12 FLOPs. When a candidate fails, its recorded partial cost is included and ≥ marks a lower bound; unrecorded interrupted work and historical retries are not imputed. Joint cost is reported separately.',
        'Intrans. is mean intransigence over all task stages: at stage t, subtract the continual model accuracy on the newly introduced task from the paired joint model accuracy on that same task. All LR candidates use the existing joint reference for the same method, order and seed, trained with the selected LR. Negative values are retained. Joint rows show a dash because joint is the reference. Values are fractions, not percentage points.',
        'Parameter and other storage columns are measured at the final stage. MiB = bytes / 1048576; KiB = bytes / 1024. Joint auxiliary storage contains model buffers and optimizer state. Joint rows retain the source method whose selected LR determined that reference; they are not a complete independent three-LR joint grid.',
        'The CSV includes each successful candidate once, plus paired joint references. test_accuracy is sample-weighted accuracy on all seen test classes at that stage. pertaskflops is task-local core plus auxiliary, populated only for TagFex. Parameter bytes are reconstructed from registered FP32 model shapes at each stage and checked against the final stored total. corrected equals the original byte count. Joint run IDs include the reference method.',
    ]
    for value in definitions:doc.add_paragraph(value)
    if data['failures']:
        doc.add_paragraph('Failed candidates','Heading 2')
        for row in data['failures']:
            doc.add_paragraph(f'{DISPLAY[row["method"]]}  O={row["order"]} S={row["seed"]} LR={row["lr"]}: {row["error"]}')
    stream=BytesIO()
    doc.save(stream)
    # Only the body and explicitly adjusted text styles are editable slots.
    # Retain all other template package parts and relationships byte-for-byte.
    with ZipFile(template) as source, ZipFile(BytesIO(stream.getvalue())) as built, ZipFile(output,'w') as target:
        for info in source.infolist():
            content=built.read(info.filename) if info.filename in ('word/document.xml','word/styles.xml') else source.read(info.filename)
            target.writestr(info,content)


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project-root',type=Path,default=Path(__file__).resolve().parent)
    parser.add_argument('--results-root',type=Path)
    parser.add_argument('--template',type=Path,required=True)
    parser.add_argument('--output-dir',type=Path,default=Path('pp2_docx_csv'))
    parser.add_argument('--experiments',nargs='+',default=['spike=spike-all','uwave=uwave-all','texture=texture-all-4'])
    args=parser.parse_args(argv)
    project=args.project_root.resolve();root=(args.results_root or project).resolve();out=args.output_dir.resolve()
    registry=runpy.run_path(str(project/'cil_experiments/order_seed_registry.py'))
    search_config=runpy.run_path(str(project/'cil_experiments/search_config.py'))
    digest=hashlib.sha256(args.template.read_bytes()).hexdigest()
    reports={}
    for spec in args.experiments:
        dataset,sep,exp=spec.partition('=')
        require(sep and dataset in registry['ORDERS_BY_DATASET'] and dataset not in reports,'Invalid or duplicate dataset')
        require(exp and all(c.isalnum() or c in '-_.' for c in exp) and exp not in ('.','..'),'Invalid experiment')
        choices=[p for p in (root,root/exp) if (p/f'search_result_{exp}').is_dir()]
        require(len(choices)==1,f'Missing/ambiguous result root: {exp}')
        for prefix in ('search_result','joint_result'):
            require(not out.is_relative_to(choices[0]/f'{prefix}_{exp}'),'Output overlaps inputs')
        require(args.template.resolve()!=out/f'{dataset}.docx','Output would overwrite template')
        reports[dataset]=load_dataset(choices[0],dataset,exp,registry['ORDERS_BY_DATASET'][dataset],
                registry['SEEDS_BY_DATASET'][dataset],search_config['LR_CANDIDATES'],search_config['PIPELINE_METHODS'])
    # Build all deliverables in a staging directory before publishing any report.
    out.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.pp2-docx-',dir=out.parent) as temporary:
        staged=Path(temporary)
        for dataset,data in reports.items():
            with (staged/f'{dataset}.csv').open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.DictWriter(stream,fieldnames=FIELDS);writer.writeheader();writer.writerows(data['csv_rows'])
            write_docx(args.template,staged/f'{dataset}.docx',data)
        require(hashlib.sha256(args.template.read_bytes()).hexdigest()==digest,'Template changed')
        out.mkdir(parents=True,exist_ok=True)
        # TemporaryDirectory has private ACLs on Windows. Copy into a new file
        # under the output directory before replacement so Word inherits the
        # destination directory's normal read permissions.
        for path in staged.iterdir():
            destination = out / path.name
            publication = out / f'.{path.name}.{uuid.uuid4().hex}.tmp'
            try:
                with publication.open('xb') as stream:
                    stream.write(path.read_bytes())
                os.replace(publication, destination)
            finally:
                publication.unlink(missing_ok=True)
    for dataset,data in reports.items():
        print(f'EXPORTED {dataset} tables={len(make_sections(data))} csv_rows={len(data["csv_rows"])} failed={len(data["failures"])}')
    return 0


if __name__=='__main__':
    raise SystemExit(main())
